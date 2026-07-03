#!/usr/bin/env python3
"""Extract Olready Analysis xlsx sheets to JSON on stdout."""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl required: pip install openpyxl", file=sys.stderr)
    sys.exit(1)


def cell_str(val) -> str:
    if val is None:
        return ""
    if isinstance(val, float) and val == int(val):
        return str(int(val))
    if isinstance(val, (datetime, date)):
        return val.isoformat()
    return str(val).strip()


def is_formula_row(row) -> bool:
    if not row:
        return True
    first = row[0]
    return isinstance(first, str) and first.startswith("=")


def read_users(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [cell_str(c).lower() for c in rows[0]]
    out = []
    for row in rows[1:]:
        if not row or not any(row):
            continue
        rec = {header[i]: cell_str(row[i]) if i < len(row) else "" for i in range(len(header))}
        email = rec.get("email", "")
        if not email or "@" not in email:
            continue
        out.append(rec)
    return out


def read_raw_data(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "RAW DATA" not in wb.sheetnames:
        return []
    ws = wb["RAW DATA"]
    rows = list(ws.iter_rows(values_only=True))
    # Row 0 = banner, row 1 = headers
    if len(rows) < 3:
        return []
    header = [cell_str(c) for c in rows[1]]
    key_map = {h.lower(): i for i, h in enumerate(header)}
    out = []
    for row in rows[2:]:
        if is_formula_row(row):
            continue
        name = cell_str(row[key_map.get("mua name", 0)]) if "mua name" in key_map else ""
        if not name or name.lower().startswith("olready"):
            continue
        out.append(
            {
                "name": name,
                "phone": cell_str(row[key_map.get("phone number", 1)]) if "phone number" in key_map else "",
                "status": cell_str(row[key_map.get("status", 2)]) if "status" in key_map else "",
                "nextFollowUp": cell_str(row[key_map.get("next follow-up", 3)]) if "next follow-up" in key_map else "",
                "city": cell_str(row[key_map.get("city", 4)]) if "city" in key_map else "",
                "source": cell_str(row[key_map.get("source", 5)]) if "source" in key_map else "",
                "userType": cell_str(row[key_map.get("user type", 6)]) if "user type" in key_map else "",
                "planInterest": cell_str(row[key_map.get("plan interest", 7)]) if "plan interest" in key_map else "",
                "notes": cell_str(row[key_map.get("notes / remarks", 8)]) if "notes / remarks" in key_map else "",
                "email": cell_str(row[key_map.get("email", 9)]) if "email" in key_map else "",
            }
        )
    return out


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=["users", "sales-raw"])
    parser.add_argument("path")
    args = parser.parse_args()
    path = Path(args.path)
    if args.mode == "users":
        print(json.dumps(read_users(path)))
    else:
        print(json.dumps(read_raw_data(path)))


if __name__ == "__main__":
    main()
